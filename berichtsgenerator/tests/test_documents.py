import io

import docx
from docx.oxml import parse_xml

from app import documents

NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' \
     'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'


def _formular(filled: bool) -> bytes:
    """Formular mit drei Arten von Kontrollkästchen und einer Tabelle."""
    d = docx.Document()
    d.add_paragraph("Standortbericht Formular")
    d.add_paragraph("Lesen: Stand der Entwicklung")
    if filled:
        d.add_paragraph("Liest kurze Texte mit Silbenbögen.")
    # 1. Inhaltssteuerelement-Kontrollkästchen (Word 2010+), Zeichen steht im Text
    box = "☒" if filled else "☐"
    p = d.add_paragraph()
    p._p.append(parse_xml(
        f'<w:sdt {NS}><w:sdtPr><w14:checkbox><w14:checked w14:val="{1 if filled else 0}"/></w14:checkbox></w:sdtPr>'
        f'<w:sdtContent><w:r><w:t>{box}</w:t></w:r></w:sdtContent></w:sdt>'))
    p.add_run(" Förderziel erreicht")
    # 2. Altes Formular-Kontrollkästchen (FORMCHECKBOX)
    p = d.add_paragraph()
    checked = "<w:checked/>" if filled else ""
    p._p.append(parse_xml(
        f'<w:r {NS}><w:fldChar w:fldCharType="begin"><w:ffData><w:checkBox><w:sizeAuto/>'
        f'<w:default w:val="0"/>{checked}</w:checkBox></w:ffData></w:fldChar></w:r>'))
    p.add_run(" Logopädie empfohlen")
    # 3. Wingdings-Symbol
    p = d.add_paragraph()
    char = "F0FE" if filled else "F0A8"
    p._p.append(parse_xml(f'<w:r {NS}><w:sym w:font="Wingdings" w:char="{char}"/></w:r>'))
    p.add_run(" Nachteilsausgleich")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Bereich", "Beobachtung"
    t.cell(1, 0).text = "Mathematik"
    t.cell(1, 1).text = "Addiert bis 100 sicher" if filled else ""
    d.add_paragraph("Unterschrift Lehrperson")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_checkboxes_and_tables():
    text = documents.extract_text("f.docx", _formular(True))
    assert "[x] Förderziel erreicht" in text
    assert "[x] Logopädie empfohlen" in text
    assert "[x] Nachteilsausgleich" in text
    assert "Mathematik | Addiert bis 100 sicher" in text
    empty = documents.extract_text("f.docx", _formular(False))
    assert "[ ] Förderziel erreicht" in empty and "[ ] Logopädie empfohlen" in empty
    assert "[ ] Nachteilsausgleich" in empty


def test_template_delta():
    template = documents.extract_text("leer.docx", _formular(False))
    filled = documents.extract_text("kind.docx", _formular(True))
    tmpl = [{"id": 7, "title": "Formular SB", "content_text": template}]
    reduced, used = documents.apply_template(filled, tmpl, "auto")
    assert used["id"] == 7
    assert "Liest kurze Texte mit Silbenbögen." in reduced
    assert "Lesen: Stand der Entwicklung" in reduced  # Überschrift als Kontext
    assert "[x] Förderziel erreicht" in reduced
    assert "Mathematik | Addiert bis 100 sicher" in reduced
    assert "Unterschrift Lehrperson" not in reduced  # unveränderter Vorlagentext entfällt
    assert "[ ]" not in reduced
    # unpassende Vorlage wird nicht automatisch verwendet
    other = [{"id": 8, "title": "Anderes", "content_text": "\n".join(f"Zeile {i} ganz anders" for i in range(10))}]
    assert documents.apply_template(filled, other, "auto") == (filled, None)
    assert documents.apply_template(filled, tmpl, "none") == (filled, None)


def test_pdf_form_fields():
    class Reader:
        def get_fields(self):
            return {
                "Name": {"/FT": "/Tx", "/V": "Lina"},
                "Leer": {"/FT": "/Tx", "/V": ""},
                "Logopädie": {"/FT": "/Btn", "/V": "/Yes"},
                "Ergotherapie": {"/FT": "/Btn", "/V": "/Off"},
            }

    assert documents._pdf_form_fields(Reader()) == ["Name: Lina", "[x] Logopädie"]


def test_template_with_repeated_placeholders():
    """Identische Platzhaltertexte dürfen ausgefüllte Felder nicht dem falschen Feld zuordnen."""
    filler = "Lorem ipsum dolor sit amet"
    fields = ["Allgemeine Anmerkungen", "Mathematik", "Deutsch", "Musik", "Personale Kompetenzen", "Ausblick"]
    template = "\n".join(f"{f}\n{filler}" for f in fields)
    filled = "\n".join(
        f"{f}\n" + ("Übernimmt Verantwortung für das Material." if f == "Personale Kompetenzen" else "")
        for f in fields
    )
    reduced = documents.strip_template(filled, template)
    assert reduced == "Personale Kompetenzen\nÜbernimmt Verantwortung für das Material."

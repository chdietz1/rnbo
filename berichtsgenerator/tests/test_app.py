import io
import os
import re
import tempfile

os.environ["BG_DATA_DIR"] = tempfile.mkdtemp()
os.environ["BG_LLM_BACKEND"] = "mock"

import docx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import documents  # noqa: E402
from app.main import app  # noqa: E402


def csrf(client, url):
    html = client.get(url).text
    return re.search(r'name="csrf" content="([^"]+)"', html).group(1)


def login(client, username, password):
    token = csrf(client, "/login")
    return client.post("/login", data={"csrf": token, "username": username, "password": password})


def test_full_flow():
    with TestClient(app) as c:
        # Ersteinrichtung
        assert c.get("/", follow_redirects=False).headers["location"] == "/login"
        assert c.get("/login", follow_redirects=False).headers["location"] == "/setup"
        token = csrf(c, "/setup")
        r = c.post("/setup", data={"csrf": token, "display_name": "Admin", "username": "admin",
                                    "password": "geheim12345", "password2": "geheim12345"})
        assert "Übersicht" in r.text

        # CSRF ist Pflicht
        assert c.post("/students", data={"first_name": "A", "last_name": "B"}).status_code == 403

        # Lehrperson anlegen
        token = csrf(c, "/admin/users")
        c.post("/admin/users", data={"csrf": token, "display_name": "Lehrerin", "username": "lp",
                                      "password": "passwort1234"})

        # Dossier + Dokument
        token = csrf(c, "/students")
        r = c.post("/students", data={"csrf": token, "first_name": "Mia", "last_name": "Muster",
                                       "class_name": "Unterstufe"})
        sid = int(r.url.path.rsplit("/", 1)[1])
        d = docx.Document()
        d.add_paragraph("Mia liest einfache Wörter. Förderziel: Silben lesen.")
        buf = io.BytesIO()
        d.save(buf)
        token = csrf(c, f"/students/{sid}")
        r = c.post(f"/students/{sid}/documents",
                   data={"csrf": token, "title": "Bericht 2024", "doc_type": "Beurteilungsbericht"},
                   files={"file": ("bericht.docx", buf.getvalue(),
                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert "gespeichert" in r.text

        # Wissensbasis
        token = csrf(c, "/knowledge")
        c.post("/knowledge", data={"csrf": token, "title": "Leitfaden", "doc_type": "Leitfaden",
                                    "text": "Formative Beurteilung Förderziel Lesen beschreibt Lernfortschritte."})

        # Generieren (Mock)
        html = c.get(f"/students/{sid}/generate").text
        token = re.search(r'name="csrf" content="([^"]+)"', html).group(1)
        rt_id = int(re.search(r'<option value="(\d+)"', html).group(1))
        doc_id = int(re.search(r'name="document_ids" value="(\d+)"', html).group(1))
        r = c.post("/api/generate", headers={"X-CSRF-Token": token},
                   json={"student_id": sid, "report_type_id": rt_id, "document_ids": [doc_id],
                         "observations": "Förderziel Lesen erreicht", "use_kb": True})
        assert r.status_code == 200
        assert "Enthält frühere Unterlagen: ja" in r.text
        assert "Enthält Grundlagen: ja" in r.text

        # Speichern, finalisieren, exportieren
        r = c.post(f"/students/{sid}/reports", data={"csrf": token, "report_type_id": rt_id,
                                                      "content": "## Lernen\n- Mia liest Silben."})
        rid = int(r.url.path.rsplit("/", 1)[1])
        r = c.post(f"/reports/{rid}", data={"csrf": token, "title": "Bericht", "period": "2025",
                                            "content": "## Lernen\n- Mia liest Silben.", "finalize": "1"})
        assert "finalisiert" in r.text
        r = c.get(f"/reports/{rid}/docx")
        assert r.status_code == 200 and r.content[:2] == b"PK"

        # Lehrperson ohne Freigabe sieht das Dossier nicht
        c.post("/logout", data={"csrf": token})
        login(c, "lp", "passwort1234")
        assert c.get(f"/students/{sid}").status_code == 404
        assert c.get(f"/reports/{rid}").status_code == 404
        html = c.get("/").text
        token = re.search(r'name="csrf" content="([^"]+)"', html).group(1)
        r = c.post("/api/generate", headers={"X-CSRF-Token": token},
                   json={"student_id": sid, "report_type_id": rt_id})
        assert r.status_code == 404


def test_chunking():
    text = "\n\n".join(f"Absatz {i}. " + "Wort " * 150 for i in range(10))
    chunks = documents.chunk_text(text)
    assert len(chunks) > 3
    assert all(len(ch) <= documents.CHUNK_SIZE + 10 for ch in chunks)


def test_bulk_import():
    import openpyxl

    with TestClient(app) as c:
        login(c, "admin", "geheim12345")
        token = csrf(c, "/admin/import")

        # Excel mit Jahrgang als Zahl und Zugriff für «lp»
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Vorname", "Nachname", "Jahrgang", "Klasse", "Lehrpersonen"])
        ws.append(["Noah", "Import", 2018, "U1", "lp, unbekannt"])
        ws.append(["", "OhneVorname", 2017, "U1", ""])
        buf = io.BytesIO()
        wb.save(buf)
        r = c.post("/admin/import", data={"csrf": token},
                   files={"file": ("liste.xlsx", buf.getvalue(), "application/octet-stream")})
        assert "1 Dossier(s) angelegt" in r.text
        assert "Zeile 3" in r.text and "unbekannt" in r.text

        # CSV (Excel-Export mit Semikolon, Windows-Kodierung), Noah ist ein Duplikat
        csv_data = "Vorname;Nachname;Jahrgang;Klasse\nNoah;Import;2018;U1\nZoé;Müller;2016;M1\n".encode("cp1252")
        r = c.post("/admin/import", data={"csrf": token},
                   files={"file": ("liste.csv", csv_data, "text/csv")})
        assert "1 Dossier(s) angelegt" in r.text and "1 bereits vorhanden" in r.text
        assert "Müller" in c.get("/students").text

        assert c.get("/admin/import/vorlage.csv").status_code == 200

        # Lehrperson «lp» sieht Noah, aber nicht Zoé, und darf nicht importieren
        c.post("/logout", data={"csrf": token})
        login(c, "lp", "passwort1234")
        page = c.get("/students").text
        assert "Import Noah" in page and "Müller" not in page
        assert c.get("/admin/import").status_code == 403


def test_skills_and_template_upload():
    from tests.test_documents import _formular

    with TestClient(app) as c:
        login(c, "admin", "geheim12345")
        token = csrf(c, "/knowledge")
        # leere Vorlage hochladen
        c.post("/knowledge", data={"csrf": token, "title": "Formular SB", "doc_type": documents.TEMPLATE_TYPE},
               files={"file": ("leer.docx", _formular(False), "application/octet-stream")})
        token = csrf(c, "/students")
        r = c.post("/students", data={"csrf": token, "first_name": "Ben", "last_name": "Vorlage"})
        sid = int(r.url.path.rsplit("/", 1)[1])
        r = c.post(f"/students/{sid}/documents", data={"csrf": token, "title": "SB 2025", "doc_type": "Standortbericht"},
                   files={"file": ("kind.docx", _formular(True), "application/octet-stream")})
        assert "Abgleich mit Vorlage" in r.text and "Vorlage abgeglichen" in r.text
        did = int(re.search(r'href="/documents/(\d+)">SB 2025', r.text).group(1))
        page = c.get(f"/documents/{did}").text
        assert "Vollständiger Originaltext" in page
        r = c.post(f"/documents/{did}/template", data={"csrf": token, "template": "none"})
        assert "Vollständiger Text wird verwendet" in r.text

        # Skills: über die Oberfläche und über die API speichern, im Generator sichtbar
        c.post("/skills", data={"csrf": token, "name": "Kommunikation", "text": "Schwerpunkt Kommunikation."})
        r = c.post("/api/skills", headers={"X-CSRF-Token": token},
                   json={"name": "Kurz", "text": "Maximal eine Seite.", "shared": True})
        assert r.status_code == 200
        page = c.get(f"/students/{sid}/generate").text
        assert "Kommunikation" in page and 'data-text="Maximal eine Seite."' in page

        # Lehrperson sieht nur den geteilten Skill und kann ihn nicht ändern
        c.post("/logout", data={"csrf": token})
        login(c, "lp", "passwort1234")
        page = c.get("/skills").text
        assert "<strong>Kurz</strong>" in page and "<strong>Kommunikation</strong>" not in page
        token = csrf(c, "/skills")
        skill_id = r.json()["id"]
        assert c.post("/skills", data={"csrf": token, "id": skill_id, "action": "delete"}).status_code == 404

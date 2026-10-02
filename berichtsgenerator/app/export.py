"""Export eines Berichts als Word-Datei (.docx)."""
from __future__ import annotations

import io
import re

import docx
from docx.shared import Pt


def report_to_docx(report: dict, student: dict, author: str, school: str) -> bytes:
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    d.add_heading(report["title"], level=0)
    meta = d.add_paragraph()
    meta.add_run(f"{school}\n").bold = True
    meta.add_run(f"Name: {student['first_name']} {student['last_name']}\n")
    if student.get("class_name"):
        meta.add_run(f"Klasse/Stufe: {student['class_name']}\n")
    if report.get("period"):
        meta.add_run(f"Periode/Datum: {report['period']}\n")
    meta.add_run(f"Verfasst von: {author}")

    for raw in report["content"].splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        if line.startswith("### "):
            d.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            d.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            d.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*•] ", line):
            d.add_paragraph(re.sub(r"^\s*[-*•] ", "", line), style="List Bullet")
        else:
            d.add_paragraph(line)

    d.add_paragraph()
    d.add_paragraph("Ort, Datum: ______________________    Unterschrift: ______________________")

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

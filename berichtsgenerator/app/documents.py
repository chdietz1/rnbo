"""Textextraktion aus hochgeladenen Dateien, Aufteilung in Abschnitte und Indexierung."""
from __future__ import annotations

import io
import re
import struct

from . import llm

SUPPORTED = (".pdf", ".docx", ".txt", ".md", ".rtf")

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200


class ExtractionError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:  # pypdf wirft diverse Fehlertypen
            raise ExtractionError(f"PDF konnte nicht gelesen werden: {e}") from e
        if not text.strip():
            raise ExtractionError(
                "Im PDF wurde kein Text gefunden (vermutlich ein Scan). "
                "Bitte eine Textversion (Word/PDF mit Text) hochladen oder den Text einfügen."
            )
    elif name.endswith(".docx"):
        import docx

        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise ExtractionError(f"Word-Datei konnte nicht gelesen werden: {e}") from e
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                cells = []
                for c in row.cells:
                    if c.text.strip() and c.text.strip() not in cells:
                        cells.append(c.text.strip())
                if cells:
                    lines.append(" | ".join(cells))
        text = "\n".join(lines)
    elif name.endswith(".rtf"):
        text = _strip_rtf(data.decode("latin-1", errors="replace"))
    elif name.endswith((".txt", ".md")):
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    else:
        raise ExtractionError(
            "Dateityp nicht unterstützt. Erlaubt: " + ", ".join(SUPPORTED)
        )
    return normalize(text)


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00ad", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _strip_rtf(rtf: str) -> str:
    rtf = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1252", "replace"), rtf)
    rtf = re.sub(r"\\par[d]?", "\n", rtf)
    rtf = re.sub(r"\{\\\*[^{}]*\}", "", rtf)
    rtf = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", rtf)
    return rtf.replace("{", "").replace("}", "")


def chunk_text(text: str) -> list[str]:
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for p in paragraphs:
        while len(p) > CHUNK_SIZE:
            if current:
                chunks.append(current)
                current = ""
            cut = p.rfind(". ", 0, CHUNK_SIZE)
            cut = cut + 1 if cut > CHUNK_SIZE // 2 else CHUNK_SIZE
            chunks.append(p[:cut].strip())
            p = p[max(0, cut - CHUNK_OVERLAP):].strip()
        if len(current) + len(p) + 2 > CHUNK_SIZE and current:
            chunks.append(current)
            tail = current[-CHUNK_OVERLAP:]
            current = tail[tail.find(" ") + 1:] + "\n\n" + p if " " in tail else p
        else:
            current = f"{current}\n\n{p}" if current else p
    if current:
        chunks.append(current)
    return chunks


def pack_embedding(vec: list[float]) -> bytes:
    return struct.pack(f"{len(vec)}f", *vec)


def unpack_embedding(blob: bytes) -> list[float]:
    return list(struct.unpack(f"{len(blob) // 4}f", blob))


def index_document(conn, document_id: int, text: str) -> None:
    conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
    chunks = chunk_text(text)
    embeddings = llm.embed(chunks) if chunks else None
    for i, chunk in enumerate(chunks):
        emb = pack_embedding(embeddings[i]) if embeddings and i < len(embeddings) else None
        conn.execute(
            "INSERT INTO chunks (document_id, idx, text, embedding) VALUES (?, ?, ?, ?)",
            (document_id, i, chunk, emb),
        )
